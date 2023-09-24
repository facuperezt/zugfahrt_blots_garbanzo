#%%
import matplotlib.pyplot as plt
import os
import pandas as pd
import numpy as np
from typing import Tuple, List
import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_interval
import pandas as pd
from openpyxl.utils.cell import coordinate_from_string as cfs


def convert_rng_to_df(tbl_coords, sheet):
    col_start = cfs(tbl_coords.split(':')[0])[0]
    col_end = cfs(tbl_coords.split(':')[1])[0]

    data_rows = []
    for row in sheet[tbl_coords]:
        data_rows.append([cell.value for cell in row])

    return pd.DataFrame(data_rows, columns=get_column_interval(col_start, col_end))

def concat_tables(ws, lysat_nr):
    ### Dictionary to hold the dfs for each table
    df_dict = {}

    ### Get the table coordinates from the worksheet table dictionary
    for tblname, tblcoord in ws.tables.items():
        # print(f'Table Name: {tblname}, Coordinate: {tblcoord}')
        date = tblname[len("table"):]
        df = convert_rng_to_df(tblcoord, ws)  # Convert to dataframe
        df.columns = df.iloc[0]
        df["Kern or Cyt"] = df.columns[0].split(' ')[0]
        df["Date"] = tblname.split('_')[0][-4:]
        df["ID"] = tblname.split('_')[1]
        df["LysatNr"] = lysat_nr
        df = df.drop(df.index[0])
        # if len(df.columns[0].split(' ')) > 1:
        #     name, _id = df.columns[0].split(' ')
        #     df["ID"] = int(_id)
        df.rename(columns= {df.columns[0] : "Experiment"}, inplace=True)

        df_dict[date] = df

    return pd.concat(
        df_dict.values()
    )


def get_experiment_and_control_pairs(df: pd.DataFrame) -> List[Tuple[List[str], str]]:
    all_exps = df["Experiment"]
    controls = [exp for exp in all_exps if exp.startswith("K ")]
    exps = [exp for exp in all_exps if exp not in controls]
    exps_ctrl_dict = {ctrl: [] for ctrl in controls}
    different_rpmi_controls_flag = True if len([ctrl for ctrl in controls if "RPMI" in ctrl]) > 1 else False
    for exp in all_exps:
        if "M1" not in exp and "RPMI" not in exp:
            exps_ctrl_dict["K DL"].append(exp)
            continue
        if not different_rpmi_controls_flag:
            exps_ctrl_dict[[ctrl for ctrl in controls if "RPMI" in ctrl][0]].append(exp)
            continue
        for time in ["24", "72"]:
            if time not in exp: continue
            exps_ctrl_dict[[ctrl for ctrl in controls if "RPMI" in ctrl and time in ctrl][0]].append(exp)
            continue
    
    return (exps_ctrl_dict.keys(), exps_ctrl_dict.values()) 

        



def apply_normalize_by_K_DL(df: pd.DataFrame, numeric_cols: List[str]) -> pd.DataFrame:
    exps_control, experiments = get_experiment_and_control_pairs(df)
    for exp, exp_ctrl in zip(experiments, exps_control):
        df.loc[df["Experiment"].isin(exp), numeric_cols] = df.loc[df["Experiment"].isin(exp), numeric_cols].divide(
            df.loc[df["Experiment"] == exp_ctrl, numeric_cols].values,
            axis= 1,
            )
    return df
    

def concatenate_all_sheets(wb) -> pd.DataFrame:
    df = pd.concat(
        [concat_tables(wb[lysat], int(lysat.split(' ')[1])) for lysat in ["Lysat "+ str(ind) for ind in range(1,5)]]
    )
    df.index = range(len(df))
    return df


def normalize_concatenated_df(df: pd.DataFrame) -> pd.DataFrame:
    non_numeric_cols = ['Date', 'ID', 'Kern or Cyt', 'Experiment', 'LysatNr']
    numeric_cols = df.columns.difference(non_numeric_cols)
    df.loc[:, numeric_cols] = df.loc[:, numeric_cols].apply(
        lambda row: row/row["H3"] if not pd.isna(row["H3"]) else row/row["GAPDH"],
        axis= 1,
        )  # Normalize by load control
    df = df.groupby(non_numeric_cols[:-2]).apply(
        apply_normalize_by_K_DL,
        numeric_cols=numeric_cols,
        )  # Normalize by K DL
    return df


def get_p_over_normal_ratio(df: pd.DataFrame) -> pd.DataFrame:
    get_ratio_from = ["Erk", "Akt", "Stat3"]
    for protein in get_ratio_from:
        df[f"p{protein}/{protein}"] = df.loc[:, f"p{protein}"].mean()/df.loc[:, protein].mean()

    return df


if __name__ == "__main__":
    filename = 'data/Lysate_WesternBlot.xlsx'
    wb = load_workbook(filename)
    df = concatenate_all_sheets(wb)
    norm_df = normalize_concatenated_df(df)
    norm_df_with_activations = norm_df.reset_index(drop=True).groupby(["LysatNr", "Kern or Cyt", "Experiment", "Date"]).apply(get_p_over_normal_ratio)

    # tmp = norm_df.reset_index(drop= True).groupby(["Date", "Kern or Cyt", "ID", "Experiment"])\
    # .apply(lambda df: df).drop(columns=["Date", "Kern or Cyt", "ID", "Experiment"]).fillna(0).to_numpy()
    # plt.imshow(tmp)

    group_per_exp = norm_df_with_activations.reset_index(drop= True).drop(columns=["Date", "ID"]).groupby(["Kern or Cyt", "Experiment"], group_keys=True)
    avg_per_exp = group_per_exp.mean()
    std_per_exp = group_per_exp.std()
    std_per_exp.drop([ind for ind in avg_per_exp.index if ind[1].startswith("K")], inplace=True)
    avg_per_exp.drop([ind for ind in avg_per_exp.index if ind[1].startswith("K")], inplace=True)
    avg_for_plot = avg_per_exp.transpose()
    std_for_plot = std_per_exp.transpose()
    avg_for_plot = avg_for_plot.reindex(["Claudin-1", "GAPDH", "H3", "Erk", "Akt", "Stat3", "pErk", "pAkt", "pStat3", "pErk/Erk", "pAkt/Akt", "pStat3/Stat3"])
    std_for_plot = std_for_plot.reindex(["Claudin-1", "GAPDH", "H3", "Erk", "Akt", "Stat3", "pErk", "pAkt", "pStat3", "pErk/Erk", "pAkt/Akt", "pStat3/Stat3"])
    sorted_idx = avg_per_exp.values.argsort(axis=0)
    out_index_name = pd.DataFrame(avg_per_exp.index.to_numpy()[sorted_idx], columns=avg_per_exp.columns)
    n = 3
    div, mod = divmod(len(avg_for_plot), n)
    for kern_or_cyt in ["Kern", "Cytosol"]:
        fig, axs = plt.subplots(div+mod, n, figsize=(12,10), sharex= True, sharey= False)
        fig.suptitle(kern_or_cyt)
        for (idx, mean), (idx, std), ax in zip(avg_for_plot.iterrows(), std_for_plot.iterrows(), axs.flatten()):
            ax: plt.Axes
            y_vals = mean.loc[[kern_or_cyt]]
            x_vals = [ind[1] for ind in y_vals.index]
            y_err = std.loc[[kern_or_cyt]]
            ax.bar(x_vals, y_vals)
            ax.errorbar(x_vals, y_vals, y_err, fmt='x', barsabove= True, ecolor="red", capsize=4)
            ax.set_xticks(x_vals, x_vals, rotation=45)
            ax.set_title(idx)
    plt.show()


# %%
